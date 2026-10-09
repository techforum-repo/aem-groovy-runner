"""Batch runs: discover roots, run once per root, combine results."""
from __future__ import annotations

import time
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from groovy_runner import database, runner, settings_store
from groovy_runner.clients.mock import MockGroovyConsoleClient
from groovy_runner.config import settings
from groovy_runner.formatters import base
from groovy_runner.scripts_registry import discover

APP = str(Path(__file__).resolve().parent.parent / "app.py")


@pytest.fixture()
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(database, "DB_PATH", tmp_path / "t.db")
    monkeypatch.setattr(settings, "output_dir", str(tmp_path / "out"))
    settings_store._cache = None
    database.initialize()
    yield tmp_path
    settings_store._cache = None


def _script(sid):
    return {s.id: s for s in discover()}[sid]


def test_batch_discovers_sites_and_runs_one_per_root(isolated):
    script = _script("page-report")
    found = []
    results = runner.run_script(MockGroovyConsoleClient(), script, {"rootPath": ["/content"]}, session_id="S",
                                batch={"excludes": list(script.system_excludes.paths), "levels": 1},
                                on_discovered=found.extend)
    assert [r.label for r in results] == ["/content/acme", "/content/acme-careers", "/content/acme-corporate"]
    assert {tuple(r.groups) for r in results} == {("/content",)}  # each part knows its entered path
    assert all(r.ok for r in results)
    assert "/content/campaigns" in found[0].skipped  # a default skip (an AEM system area that can be a page)
    assert "/content/dam" not in [r.label for r in results]  # not a page, so never a site: no exclude needed
    assert "run.batch_discovered" in [e["action"] for e in database.list_audit()]


def test_without_batch_runs_the_entered_path_itself(isolated):
    results = runner.run_script(MockGroovyConsoleClient(), _script("page-report"), {"rootPath": ["/content"]}, session_id="S")
    assert [r.label for r in results] == ["/content"]


def test_discovery_failure_is_reported_as_a_row(isolated):
    class Broken(MockGroovyConsoleClient):
        def run_script(self, script):
            raise RuntimeError("AEM returned HTTP 500: boom")
    results = runner.run_script(Broken(), _script("assets-by-type"), {"rootPath": ["/content/dam"]}, session_id="S",
                                batch={"excludes": [], "levels": 1})
    assert [r.label for r in results] == ["/content/dam (discovery)"] and "boom" in results[0].error


def test_combine_concatenates_rows_into_one_file(isolated):
    script = _script("page-report")
    results = runner.run_script(MockGroovyConsoleClient(), script, {"rootPath": ["/content/a", "/content/b"]},
                                session_id="S")
    out = runner.format_combined(results, "generic-excel", "S")
    rows = openpyxl.load_workbook(out).active.max_row - 1
    assert rows == sum(r.row_count for r in results)
    assert "format.combined" in [e["action"] for e in database.list_audit()]
    with pytest.raises(ValueError, match="at least two"):
        runner.format_combined(results[:1], "generic-excel", "S")
    Path(results[0].json_path).write_text("[]")
    with pytest.raises(ValueError, match="changed on disk"):
        runner.format_combined(results, "generic-excel", "S")


def test_excel_row_limit_gives_clear_error(tmp_path, monkeypatch):
    monkeypatch.setattr(base, "EXCEL_MAX_DATA_ROWS", 3)
    with pytest.raises(ValueError, match="more than one Excel sheet can hold"):
        base.write_styled_excel(pd.DataFrame({"a": range(4)}), tmp_path / "x.xlsx", "S")


def test_ui_batch_run_and_combine(isolated):
    at = AppTest.from_file(APP, default_timeout=60).run()
    at.selectbox(key="script_id").set_value("assets-by-type").run()
    at.text_area(key="in:assets-by-type:rootPath").set_value("/content/dam").run()
    assert at.checkbox(key="skip_on:assets-by-type").value is True  # system areas skipped by default, batch or not
    assert "/content/dam/projects" in at.text_area(key="skip_list:assets-by-type").value
    at.checkbox(key="batch_on:assets-by-type").check().run()
    button = [b for b in at.button if b.label.startswith("▶️")][0]
    assert button.label == "▶️ Discover & run (1 path · mock)"
    button.click().run()
    for _ in range(40):
        if "run_job" not in at.session_state:
            break
        time.sleep(0.25)
        at.run()
    table = at.dataframe[0].value
    assert list(table["Input"]) == ["/content/dam (batch, 4 parts)"]  # ONE entry for the entered path
    assert table["Output"][0].startswith("dam_") and table["Output"][0].endswith(".generic-excel.xlsx")
    parts = at.dataframe[1].value  # the internal split, under an expander
    assert list(parts["Part"]) == ["/content/dam (direct items)", "/content/dam/acme", "/content/dam/acme-campaigns",
                                   "/content/dam/acme-legal"]
    assert any("`/content/dam` was split into 4 part(s)" in i.value for i in at.info)
    assert any("One file for `/content/dam`: **dam_" in s.value and "from 4 batch part(s)" in s.value for s in at.success)
    downloads = [b.label for b in at.get("download_button") if b.label.startswith("⬇️ dam_")]
    assert len(downloads) == 1 and downloads[0].endswith(".generic-excel.xlsx")
    assert not [b for b in at.get("download_button") if "acme" in b.label]  # no per-folder files offered


def _start_job(script, values, batch, prior=None, client=None):
    from groovy_runner import jobs
    job = jobs.start(client or MockGroovyConsoleClient(), script, values, session_id="S", formatter_id="generic-excel",
                     total=1, batch=batch, prior_results=prior)
    job.thread.join(30)
    assert job.done and job.error is None
    return job


def test_batch_gives_one_combined_file_per_entered_path(isolated):
    script = _script("page-report")
    job = _start_job(script, {"rootPath": ["/content", "/content/x"]}, {"excludes": [], "levels": 1})
    assert len(job.combined) == 2  # one per entered path, each from its 3+ discovered sites
    by_name = {Path(p).name.split("_")[0]: p for _, p in job.combined}  # same short names as other outputs
    assert {root for root, _ in job.combined} == {"/content", "/content/x"}  # keyed by entered path
    assert set(by_name) == {"content", "x"}
    content_roots = next(d.labels for d in job.discoveries if d.root == "/content")
    rows = openpyxl.load_workbook(by_name["content"]).active.max_row - 1
    assert rows == sum(r.row_count for r in job.results if r.label in content_roots)  # every root's rows, one file
    assert all(not r.outputs for r in job.results)  # no per-root Excel


def test_retry_rebuilds_the_combined_file_with_the_retried_roots(isolated):
    script = _script("page-report")

    class FlakyOnCareers(MockGroovyConsoleClient):
        def run_script(self, script_text):
            from groovy_runner.groovy_script import extract_config
            if extract_config(script_text).get("rootPath") == "/content/acme-careers":
                raise RuntimeError("AEM returned HTTP 503: busy")
            return super().run_script(script_text)

    first = _start_job(script, {"rootPath": ["/content"]}, {"excludes": [], "levels": 1}, client=FlakyOnCareers())
    assert any("1 part(s) not included" in t for _, t in first.messages)
    first_rows = openpyxl.load_workbook(first.combined[0][1]).active.max_row - 1

    retry = _start_job(script, {"rootPath": ["/content"]}, {"excludes": [], "levels": 1},
                       prior=[r for r in first.results if r.ok])
    assert not any("not included" in t for _, t in retry.messages)
    assert openpyxl.load_workbook(retry.combined[0][1]).active.max_row - 1 > first_rows


def test_too_big_to_combine_falls_back_to_one_file_per_root(isolated, monkeypatch):
    monkeypatch.setattr(base, "EXCEL_MAX_DATA_ROWS", 5)
    script = _script("page-report")

    class FewRows(MockGroovyConsoleClient):  # each root small enough alone, too many rows together
        def run_script(self, script_text):
            from groovy_runner.groovy_script import extract_config
            config = extract_config(script_text)
            if "kind" in config:
                return super().run_script(script_text)
            import json
            from groovy_runner.clients.groovy_console import GroovyResult
            from groovy_runner.groovy_script import JSON_END, JSON_START
            rows = [{"Path": f"{config['rootPath']}/p{i}"} for i in range(3)]
            return GroovyResult(output=f"{JSON_START}\n{json.dumps({'rows': rows})}\n{JSON_END}", result="",
                                exception="", running_time="1")

    job = _start_job(script, {"rootPath": ["/content"]}, {"excludes": [], "levels": 1}, client=FewRows())
    assert job.combined == []
    assert all("generic-excel" in r.outputs for r in job.results if r.ok)
    assert any("couldn't combine into one file" in t for _, t in job.messages)


def test_ui_skips_apply_without_batch_and_can_be_switched_off(isolated):
    at = AppTest.from_file(APP, default_timeout=60).run()
    at.selectbox(key="script_id").set_value("assets-by-type").run()
    at.text_area(key="in:assets-by-type:rootPath").set_value("/content/dam").run()
    at.text_area(key="in:assets-by-type:excludedFolders").set_value("/content/dam/acme/archive").run()

    def sent_config():
        from groovy_runner.groovy_script import extract_config
        code = [c.value for c in at.code if "__gr_guard" in c.value][0]  # "script exactly as sent" preview
        return extract_config(code)

    assert sent_config()["excludedFolders"][:2] == ["/content/dam/acme/archive", "/content/dam/projects"]
    at.checkbox(key="skip_on:assets-by-type").uncheck().run()
    assert sent_config()["excludedFolders"] == ["/content/dam/acme/archive"]
    at.checkbox(key="skip_on:assets-by-type").check().run()
    at.text_area(key="skip_list:assets-by-type").set_value("/content/projects").run()
    assert [b for b in at.button if b.label.startswith("▶️")][0].disabled
    assert any("Skipped system areas must start with /content/dam" in c.value for c in at.caption)


def _wait(at):
    for _ in range(60):
        if "run_job" not in at.session_state:
            return
        time.sleep(0.25)
        at.run()


def test_fresh_batch_never_borrows_results_from_an_earlier_run(isolated, monkeypatch):
    """Review fix: a new batch whose root fails must not fill the gap with an
    older result of that root (possibly run with different inputs)."""
    at = AppTest.from_file(APP, default_timeout=60).run()
    at.selectbox(key="script_id").set_value("assets-by-type").run()
    at.text_area(key="in:assets-by-type:rootPath").set_value("/content/dam").run()
    at.checkbox(key="batch_on:assets-by-type").check().run()
    [b for b in at.button if b.label.startswith("▶️")][0].click().run()
    _wait(at)
    assert any("from 4 batch part(s)" in s.value for s in at.success)

    original = MockGroovyConsoleClient.run_script

    def fail_legal(self, script_text):
        from groovy_runner.groovy_script import extract_config
        if extract_config(script_text).get("rootPath") == "/content/dam/acme-legal":
            raise RuntimeError("AEM returned HTTP 503: busy")
        return original(self, script_text)
    monkeypatch.setattr(MockGroovyConsoleClient, "run_script", fail_legal)
    at.selectbox(key="in:assets-by-type:mode").set_value("list").run()  # different inputs this time
    [b for b in at.button if b.label.startswith("▶️")][0].click().run()
    _wait(at)
    warning = [w.value for w in at.warning if "One file for `/content/dam`" in w.value]
    assert warning and "from 3 batch part(s)" in warning[0] and "1 part(s) not included" in warning[0]

    monkeypatch.setattr(MockGroovyConsoleClient, "run_script", original)
    [b for b in at.button if b.label.startswith("Retry")][0].click().run()  # same batch: may reuse its own successes
    _wait(at)
    assert any("from 4 batch part(s)" in s.value for s in at.success)
    combined = [b.label for b in at.get("download_button") if b.label.startswith("⬇️ dam_")]
    assert len(combined) == 1  # the rebuilt file replaced the partial one for /content/dam


def test_history_lists_combined_files_and_each_button_downloads_its_own_file(isolated):
    at = AppTest.from_file(APP, default_timeout=60).run()
    at.selectbox(key="script_id").set_value("page-report").run()
    at.text_area(key="in:page-report:rootPath").set_value("/content").run()
    at.checkbox(key="batch_on:page-report").check().run()
    [b for b in at.button if b.label.startswith("▶️")][0].click().run()
    _wait(at)
    at.sidebar.radio[0].set_value("History").run()
    assert not at.exception, at.exception
    labels = [b.label for b in at.get("download_button")]
    assert any(l.startswith("⬇️ content_") for l in labels)
    # Each button must serve ITS file. Capture the lazy data functions at render time, then call them later,
    # which is exactly when late-binding closures would hand back the wrong (last-assigned) file.
    import streamlit as st
    captured = []
    real = st.download_button

    def spy(label, data, file_name=None, *args, **kwargs):
        captured.append((file_name, data))
        return real(label, data, file_name, *args, **kwargs)
    st.download_button = spy
    try:
        at.run()
    finally:
        st.download_button = real
    by_name = {}
    for event in database.list_audit():
        if event["action"] == "format.combined":
            out = __import__("json").loads(event["details_json"])["output"]
            by_name[Path(out).name] = out
    for run in database.list_runs():
        for p in (run["json_path"], run["executed_script_path"]):
            if p:
                by_name[Path(p).name] = p
    assert len(captured) >= 2
    for file_name, data in captured:
        assert data() == Path(by_name[file_name]).read_bytes(), file_name


def test_new_batch_with_no_successes_removes_the_old_combined_file(isolated, monkeypatch):
    """Review fix: a later batch for the same path that produces no combined file must not leave the previous
    batch's file listed as if it were current."""
    at = AppTest.from_file(APP, default_timeout=60).run()
    at.selectbox(key="script_id").set_value("assets-by-type").run()
    at.text_area(key="in:assets-by-type:rootPath").set_value("/content/dam").run()
    at.checkbox(key="batch_on:assets-by-type").check().run()
    [b for b in at.button if b.label.startswith("▶️")][0].click().run()
    _wait(at)
    combined = lambda: [b.label for b in at.get("download_button") if b.label.startswith("⬇️ dam_")]  # noqa: E731
    assert len(combined()) == 1

    original = MockGroovyConsoleClient.run_script

    def fail_all_roots(self, script_text):
        from groovy_runner.groovy_script import extract_config
        if "kind" not in extract_config(script_text):  # discovery still works; every root run fails
            raise RuntimeError("AEM returned HTTP 503: busy")
        return original(self, script_text)
    monkeypatch.setattr(MockGroovyConsoleClient, "run_script", fail_all_roots)
    [b for b in at.button if b.label.startswith("▶️")][0].click().run()
    _wait(at)
    assert combined() == []  # the old file is no longer offered as the result for /content/dam
    assert any("No combined file for `/content/dam`: none of its 4 batch part(s) succeeded" in w.value for w in at.warning)

    monkeypatch.setattr(MockGroovyConsoleClient, "run_script", original)
    at.selectbox(key="run_formatter:assets-by-type").set_value("__none__").run()
    [b for b in at.button if b.label.startswith("▶️")][0].click().run()
    _wait(at)
    assert combined() == [] and any("no combined Excel file was made" in i.value for i in at.info)


# --- review fixes (framework-side batching) -----------------------------------

def test_overlapping_entered_paths_each_get_the_shared_parts(isolated):
    """A part found under two entered paths (/content/dam at Levels 2, /content/dam/acme at Levels 1) runs once
    but belongs to both, so each path's file and results entry include it."""
    class Overlapping(MockGroovyConsoleClient):
        def _discover(self, config):
            return {"root": config["root"], "kind": "folder", "children": ["/content/dam/acme/x"], "direct": [],
                    "skipped": []}
    script = _script("assets-by-type")
    job = _start_job(script, {"rootPath": ["/content/dam", "/content/dam/acme"]}, {"excludes": [], "levels": 1},
                     client=Overlapping())
    assert [f.label for f in job.results] == ["/content/dam/acme/x"]  # run once
    assert job.results[0].groups == ["/content/dam", "/content/dam/acme"]
    groups = runner.batch_groups(job.results)
    assert {root for (_, root, _) in groups} == {"/content/dam", "/content/dam/acme"}
    assert {root for root, _ in job.combined} == {"/content/dam", "/content/dam/acme"}


def test_unticked_include_root_is_not_reported_as_a_missing_part(isolated):
    script = _script("page-report")
    job = _start_job(script, {"rootPath": ["/content/acme"], "includeRoot": False}, {"excludes": [], "levels": 1})
    assert "/content/acme (direct items)" not in [f.label for f in job.results]
    assert not any("not included" in t for _, t in job.messages), job.messages
    assert all(f.ok for f in job.results)
    assert all(config_root_kept(f) for f in job.results)  # the parts' own roots stay in


def config_root_kept(f):
    import json
    return json.loads(database.get_run(f.run_id)["config_json"]).get("includeRoot") is True


def test_summary_rows_are_added_up_across_parts():
    from groovy_runner.scripts_registry import MergeRows
    rows = [{"Format": "application/pdf", "Assets": 2, "Total Size (MB)": 1.25},
            {"Format": "image/png", "Assets": 1, "Total Size (MB)": 0.5},
            {"Format": "application/pdf", "Assets": 3, "Total Size (MB)": 2.5}]
    merge = MergeRows(group_by=("Folder", "Format"), sum=("Assets", "Total Size (MB)"))
    assert runner.merge_rows(rows, merge) == [{"Format": "application/pdf", "Assets": 5, "Total Size (MB)": 3.75},
                                              {"Format": "image/png", "Assets": 1, "Total Size (MB)": 0.5}]
    listed = [{"Asset Path": "/a", "Format": "x"}]
    assert runner.merge_rows(listed, merge) == listed  # list-mode rows aren't totals: untouched


def test_batched_summary_gives_one_total_row_per_type(isolated):
    script = _script("assets-by-type")
    job = _start_job(script, {"rootPath": ["/content/dam"], "mode": "summary"}, {"excludes": [], "levels": 1})
    sheet = openpyxl.load_workbook(job.combined[0][1]).active
    header = [c.value for c in sheet[1]]
    formats = [row[header.index("Format")] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert len(formats) == len(set(formats))  # each type once
    total = sum(int(row[header.index("Assets")]) for row in sheet.iter_rows(min_row=2, values_only=True))
    assert total == sum(f.meta.get("matchedAssets", 0) for f in job.results if f.ok)


def test_root_relative_inputs_block_batch():
    from groovy_runner.scripts_registry import batch_conflicts
    assert batch_conflicts(_script("page-report"), {"maxDepth": 2})
    assert batch_conflicts(_script("page-report"), {"maxDepth": 0}) == []
    assert batch_conflicts(_script("assets-by-type"), {"maxAssets": 100})
    assert batch_conflicts(_script("asset-reference-report"), {}) == []


def test_ui_two_batches_of_the_same_path_keep_their_own_files(isolated):
    """Review fix: an older batch row must not show or overwrite the newer batch's file."""
    at = AppTest.from_file(APP, default_timeout=60).run()
    at.selectbox(key="script_id").set_value("assets-by-type").run()
    at.text_area(key="in:assets-by-type:rootPath").set_value("/content/dam").run()
    at.checkbox(key="batch_on:assets-by-type").check().run()
    [b for b in at.button if b.label.startswith("▶️")][0].click().run()
    _wait(at)
    time.sleep(1.1)  # a different second: a different file name
    [b for b in at.button if b.label.startswith("▶️")][0].click().run()
    _wait(at)
    outputs = list(at.dataframe[0].value["Output"])
    assert len(outputs) == 2 and outputs[0] != outputs[1] and all(o.startswith("dam_") for o in outputs)


def test_ui_nothing_found_under_a_path_is_reported(isolated, monkeypatch):
    original = MockGroovyConsoleClient._discover
    monkeypatch.setattr(MockGroovyConsoleClient, "_discover",
                        lambda self, config: {**original(self, config), "children": [], "direct": []})
    at = AppTest.from_file(APP, default_timeout=60).run()
    at.selectbox(key="script_id").set_value("page-report").run()
    at.text_area(key="in:page-report:rootPath").set_value("/content").run()
    at.checkbox(key="batch_on:page-report").check().run()
    [b for b in at.button if b.label.startswith("▶️")][0].click().run()
    _wait(at)
    assert any("Nothing to run under `/content`: no pages found" in w.value for w in at.warning)


def test_ui_batch_refused_while_a_root_relative_input_is_set(isolated):
    at = AppTest.from_file(APP, default_timeout=60).run()
    at.selectbox(key="script_id").set_value("page-report").run()
    at.text_area(key="in:page-report:rootPath").set_value("/content").run()
    at.number_input(key="in:page-report:maxDepth").set_value(2).run()
    assert not [b for b in at.button if b.label.startswith("▶️")][0].disabled  # fine without Batch
    at.checkbox(key="batch_on:page-report").check().run()
    assert [b for b in at.button if b.label.startswith("▶️")][0].disabled
